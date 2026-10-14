# -*- coding: utf-8 -*-
"""Revisão pontual e rastreável da minuta POP v1.10 de setembro de 2026.

Preserva a fonte original e o desenho do DOCX. Não substitui validação jurídica
ou aprovação institucional. Uso: python tools/revise_pop_v1101.py ORIGEM DESTINO
"""

from __future__ import annotations

import argparse
import hashlib
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


SOURCE_SHA256 = "7187a9589be1d5e83655f4c66f5fa9544d319f8a3e4bcff554d8dce927e8f35c"
FIGURE_11_SOURCE_SHA256 = "1cef96ecb34ee11ec40a527dbd2be59c79dfa3aab0f84e113e21ea0f74bbcc42"
FIGURE_11_REVISION_SHA256 = "7ad03b92d4f84c4f43a0327a5854259f852c30007085f4f3c7001b01fdb0a57f"


def replace_run(paragraph, before: str, after: str) -> None:
    matches = [run for run in paragraph.runs if before in run.text]
    if len(matches) != 1:
        raise ValueError(f"Esperada uma ocorrência de {before!r}; encontradas {len(matches)}")
    matches[0].text = matches[0].text.replace(before, after, 1)


def add_sentence(paragraph, sentence: str) -> None:
    paragraph.add_run(" " + sentence)


def set_cell_from_cloned_row(cell, value: str) -> None:
    paragraph = cell.paragraphs[0]
    if not paragraph.runs:
        paragraph.add_run(value)
    else:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""


def revise(source: Path, destination: Path) -> None:
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    if digest != SOURCE_SHA256:
        raise ValueError(f"Fonte inesperada: SHA-256 {digest}")
    if destination.exists():
        raise FileExistsError(f"Destino já existe: {destination}")

    doc = Document(source)
    if len(doc.paragraphs) != 824 or len(doc.tables) != 70:
        raise ValueError("A estrutura da fonte mudou; revisar localizadores antes de editar")

    # Capa e controle documental: revisão limitada, sem mudar a data-base
    # normativa integral de 1º de setembro de 2026.
    replace_run(doc.paragraphs[10], "Versão: 1.10", "Versão: 1.10.1")
    replace_run(doc.paragraphs[10], "Data: setembro de 2026", "Data: outubro de 2026")
    add_sentence(
        doc.tables[0].cell(9, 1).paragraphs[0],
        "Em 13 de outubro de 2026 houve revisão pontual de linguagem, do exemplo "
        "do Anexo F e da compatibilização CAR/LAC; a base normativa completa não "
        "foi reavaliada nesta revisão.",
    )

    history = doc.tables[1]
    history._tbl.append(deepcopy(history.rows[-1]._tr))
    new_row = history.rows[-1]
    for cell, value in zip(
        new_row.cells,
        (
            "1.10.1",
            "13 de outubro de 2026",
            "Revisão pontual: explicitação da tensão CAR/LAC sem solução normativa "
            "automática; substituição de expressões operacionais dispensadas; "
            "generalização de exemplo do Anexo F; revisão da rastreabilidade.",
            "Rafael Valgrande Augusto",
            "",
        ),
        strict=True,
    ):
        set_cell_from_cloned_row(cell, value)

    # O art. 9º, § 6º, da Lei Federal 15.190/2025 e o art. 64, III, do
    # Decreto Estadual 9.541/2025 devem ser confrontados. O POP não decide
    # conflito de normas nem dispensa ato exigível por interpretação própria.
    car_note = (
        "Na LAC em imóvel rural, confrontar o art. 64, III, do Decreto Estadual "
        "nº 9.541/2025 com o art. 9º, § 6º, da Lei Federal nº 15.190/2025, "
        "quando se tratar de infraestrutura de energia sem relação com as "
        "atividades agropecuárias do imóvel. Havendo aparente conflito no caso "
        "concreto, registrar os fundamentos e solicitar validação jurídica ou "
        "institucional antes de condicionar a licença ao CAR; não presumir "
        "revogação ou dispensa automática."
    )
    add_sentence(doc.paragraphs[243], car_note)
    add_sentence(doc.tables[6].cell(11, 2).paragraphs[0], car_note)

    # Linguagem operacional: não se trata de uma nova escala de gravidade.
    replacements = [
        (doc.paragraphs[412], "pontos críticos", "pontos que exigem intervenção ou acompanhamento"),
        (doc.paragraphs[506], "análise crítica", "avaliação fundamentada"),
        (doc.paragraphs[570], "pontos críticos do processo", "aspectos do processo que exigem verificação em campo"),
        (doc.tables[45].cell(5, 1).paragraphs[0], "Pontos críticos", "Pontos que exigem acompanhamento"),
        (doc.tables[63].cell(1, 1).paragraphs[0], "análise crítica", "avaliação técnica"),
        (doc.tables[63].cell(3, 1).paragraphs[0], "análise crítica", "interpretação técnica"),
        (doc.tables[5].cell(20, 1).paragraphs[0], "baixo risco", "regras de dispensa quando aplicáveis"),
        (doc.tables[51].cell(10, 0).paragraphs[0], "Atividades de baixo risco", "Critérios de dispensa da DLAM"),
        (doc.tables[51].cell(10, 2).paragraphs[0], "Classificação de baixo risco dos atos ambientais", "Classificação de atividades para fins de dispensa da declaração"),
        (doc.tables[13].cell(2, 1).paragraphs[0], "quando o caso se enquadra como baixo potencial e com compromisso do empreendedor", "quando o caso atende aos critérios normativos da modalidade e aos compromissos do empreendedor"),
        (doc.paragraphs[806], "Intervenções de baixo impacto em ", "Intervenções em "),
        (doc.paragraphs[812], "Classificação de atividades econômicas consideradas de baixo risco associada aos atos administrativos de licenciamento ambiental", "Critérios de classificação de atividades econômicas para atos de licenciamento ambiental"),
    ]
    for paragraph, before, after in replacements:
        replace_run(paragraph, before, after)

    replace_run(
        doc.paragraphs[752],
        "O exemplo de peticionamento da UHE Segredo em 2026 demonstra que",
        "Os documentos de peticionamento analisados em 2026 demonstram que",
    )

    # A Figura 11 é uma imagem raster: a busca textual no DOCX não detecta
    # "pontos críticos". Substituir somente essa figura por versão conferida.
    image_path = Path(__file__).resolve().parent / "assets" / "pop-figura-11-v1101.png"
    image_bytes = image_path.read_bytes()
    if hashlib.sha256(image_bytes).hexdigest() != FIGURE_11_REVISION_SHA256:
        raise ValueError("Figura 11 revisada inesperada")
    embed = doc.inline_shapes[10]._inline.xpath(".//a:blip")[0].get(qn("r:embed"))
    image_part = doc.part.related_parts[embed]
    if hashlib.sha256(image_part.blob).hexdigest() != FIGURE_11_SOURCE_SHA256:
        raise ValueError("Figura 11 original inesperada")
    image_part._blob = image_bytes

    core = doc.core_properties
    core.title = "POP Licenciamento de Hidrelétricas IAT v1.10.1"
    core.subject = "Minuta técnica de procedimento operacional revisada pontualmente em outubro de 2026"
    core.description = (
        "Versão 1.10.1. Revisão pontual em 13 de outubro de 2026; "
        "data-base normativa integral preservada em 1º de setembro de 2026."
    )
    core.revision = max(1, core.revision + 1)
    core.modified = datetime.now(timezone.utc)
    core.author = "Rafael Valgrande Augusto"
    core.last_modified_by = "Rafael Valgrande Augusto"

    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)
    print(f"Gerado: {destination}")
    print(f"SHA-256: {hashlib.sha256(destination.read_bytes()).hexdigest()}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    args = parser.parse_args()
    revise(args.source, args.destination)


if __name__ == "__main__":
    main()
